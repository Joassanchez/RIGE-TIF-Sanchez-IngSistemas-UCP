#!/usr/bin/env python3
"""Genera el Libro de trabajo oficial (.xlsx) desde 03-requisitos/libro/.

Copia la plantilla de la cátedra, borra las filas de ejemplo (ocre), carga las filas
desde el Markdown solo en celdas de datos (nunca en columnas con fórmula) y guarda en
la carpeta indicada con la nomenclatura AAAAMMDD_CatalogoRequisitos_Equipo_vN.xlsx.
La hoja «Recursos» se toma del Instrumento 34 (instrumentos/), su fuente única.
La hoja «Panel» se recalcula al abrir el archivo en Excel: es el control final del autor.

Uso: python tools/exportar_libro.py --destino 03-requisitos
"""
import argparse, datetime, pathlib, re, shutil, sys, unicodedata
import openpyxl
from openpyxl.styles import PatternFill

RAIZ = pathlib.Path(__file__).resolve().parent.parent
LIBRO = RAIZ / "03-requisitos" / "libro"
RECURSOS = RAIZ / "instrumentos" / "instrumento-34-recursos.md"
PLANTILLA = RAIZ / "catedra" / "originales" / "04_Libro_de_Trabajo_AE2.xlsx"
PRIMERA = 5
OCRE = "FFFDF4E2"
BLANCO = PatternFill(fill_type=None)

def norm(t):
    t = unicodedata.normalize("NFKD", re.sub(r"[*¿?]", "", t or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", t).strip()

def limpiar(v):
    v = re.sub(r"\\([\\`*_{}\[\]()#+\-.!|<>~\"'])", r"\1", (v or "").strip()).strip("*").strip()
    return "" if re.fullmatch(r"\[(DATO|DECISIÓN) PENDIENTE[^\]]*\]", v) else v

def tablas_md(texto):
    """Devuelve lista de tablas pipe: (encabezados, filas)."""
    res, act = [], []
    for ln in texto.split("\n") + [""]:
        if ln.startswith("|"): act.append(ln)
        elif act:
            filas = [[c.strip() for c in l.strip().strip("|").split("|")] for l in act if not re.match(r"^\|[-:| ]+\|$", l.strip())]
            if filas: res.append(([norm(h) for h in filas[0]], filas[1:]))
            act = []
    return res

def yaml_equipo():
    m = re.search(r'^equipo:\s*"?([^"#\n]+)', (RAIZ / "informe" / "datos-autor.yaml").read_text(encoding="utf-8"), re.M)
    return m.group(1).strip() if m else "Equipo"

def columnas_formula(ws):
    return {c.column for c in ws[PRIMERA] if isinstance(c.value, str) and c.value.startswith("=")} | \
           {c.column for c in ws[PRIMERA + 1] if isinstance(c.value, str) and c.value.startswith("=")}

def vaciar_ejemplos(ws, formulas):
    for fila in ws.iter_rows(min_row=PRIMERA, max_row=ws.max_row):
        for c in fila:
            if c.column in formulas or type(c).__name__ == "MergedCell": continue
            if c.fill and c.fill.fgColor and c.fill.fgColor.rgb == OCRE:
                c.fill = BLANCO
            c.value = None

def escribir(ws, filas):
    formulas = columnas_formula(ws); vaciar_ejemplos(ws, formulas)
    for i, fila in enumerate(filas):
        for j, v in enumerate(fila, start=1):
            if j in formulas or v in (None, ""): continue
            celda = ws.cell(row=PRIMERA + i, column=j)
            if type(celda).__name__ == "MergedCell": continue
            celda.value = v
    return len(filas)

def desde_tabla(archivo, alias, filtro=None):
    """alias: lista (una por columna del xlsx) de claves normalizadas aceptadas del encabezado Markdown."""
    filas = []
    for enc, cuerpo in tablas_md((LIBRO / archivo).read_text(encoding="utf-8")):  # una ruta absoluta (RECURSOS) ignora LIBRO
        idx = []
        for claves in alias:
            k = next((enc.index(h) for h in enc if any(h.startswith(a) for a in claves)), None) if claves else None
            idx.append(k)
        if all(k is None for k in idx) or (filtro and not filtro(enc)): continue
        for f in cuerpo:
            filas.append([limpiar(f[k]) if k is not None and k < len(f) else "" for k in idx])
    return filas

def catalogo():
    filas = []
    orden = lambda p: (p.stem.startswith("RNF"), int(re.search(r"\d+", p.stem).group()))
    for p in sorted((LIBRO / "catalogo").glob("*.md"), key=orden):
        d = {norm(a): b for a, b in re.findall(r"^\| ([^|]+?) \| (.*?) \|$", p.read_text(encoding="utf-8"), re.M)}
        g = lambda k: limpiar(d.get(norm(k), ""))
        cat = g("Categoría (si es no funcional)")
        it = g("Iteración prevista")
        filas.append([g("ID"), g("Enunciado"), g("Tipo"), "" if cat == "—" else cat, g("Prioridad"), g("Motivo de la prioridad"),
                      g("Criterio de aceptación"), g("Trazabilidad"), g("Estado de validación"),
                      "" if it == "Sin asignar" else it, g("¿Integra el MVP?")])
    return filas

def iteraciones():
    filas = []
    for f in desde_tabla("iteraciones.md", [["iteracion"], ["fechas"], ["objetivo"], ["requisitos"], ["entregable"], ["horas"]],
                         filtro=lambda enc: "fechas" in enc):
        m = re.match(r"(\S+)\s+al\s+(\S+)", f[1] or "")
        filas.append([f[0], m.group(1) if m else f[1], m.group(2) if m else "", f[2], f[3], f[4], f[5]])
    return filas

def main():
    ap = argparse.ArgumentParser(description="Genera el Libro de trabajo .xlsx.")
    ap.add_argument("--destino", required=True, help="carpeta de salida relativa a la raíz, según la consigna")
    carpeta = RAIZ / ap.parse_args().destino; carpeta.mkdir(parents=True, exist_ok=True)
    hoy = datetime.date.today().strftime("%Y%m%d"); eq = yaml_equipo(); n = 1
    while list(carpeta.glob(f"*_CatalogoRequisitos_{eq}_v{n}.xlsx")): n += 1
    destino = carpeta / f"{hoy}_CatalogoRequisitos_{eq}_v{n}.xlsx"
    shutil.copy(PLANTILLA, destino)
    wb = openpyxl.load_workbook(destino)
    informe = {
        "Catálogo": escribir(wb["Catálogo"], catalogo()),
        "Trazabilidad": escribir(wb["Trazabilidad"], desde_tabla("trazabilidad.md",
            [["cod"], ["hallazgo"], ["fuente"], ["requisitos derivados", "requisito"], ["si no deriva"]], filtro=lambda e: any(h.startswith("hallazgo") for h in e))),
        "Entidades": escribir(wb["Entidades"],
            [f[:4] + ["Entidad"] + f[5:] for f in desde_tabla("entidades.md",
                [["entidad"], [], [], [], [], ["fuente"], ["relaciones"]], filtro=lambda e: any(h.startswith("definicion") for h in e))]
            + desde_tabla("entidades.md", [["candidata"], [], [], [], ["reclasificacion"], ["criterio"], []],
                filtro=lambda e: any(h.startswith("reclasificacion") for h in e))),
        "Reglas": escribir(wb["Reglas"], desde_tabla("reglas.md",
            [["cod"], ["tipo"], ["enunciado"], ["fuente"], ["estado"], ["requisito"]])),
        "Glosario": escribir(wb["Glosario"], desde_tabla("glosario.md",
            [["termino"], ["definicion"], ["falsos amigos", "sinonimos"], ["fuente"], ["en disputa"]])),
        "Iteraciones": escribir(wb["Iteraciones"], iteraciones()),
        "Recursos": escribir(wb["Recursos"], [f for f in desde_tabla(RECURSOS,
            [["tipo", "clase"], ["recurso", "concepto"], ["cantidad"], ["unidad"], ["costo unitario"], [], ["fuente"]]) if f[1]]),
    }
    wb.save(destino)
    print(destino.relative_to(RAIZ))
    for hoja, k in informe.items(): print(f"  {hoja}: {k} filas")
    print("Abrí el archivo en Excel y revisá la hoja «Panel».")

if __name__ == "__main__":
    main()
